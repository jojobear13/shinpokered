#!/usr/bin/env python3
"""Applies the learnset audit's "Recommended Remaster Learnsets" sheet to
data/evos_moves.asm and the starting moves in data/baseStats/*.asm.

Every row is applied, CONDITIONAL ones included: the move goes to the
proposed level (taken out of wherever the species had it; a starting move
stays one only if the level is 1), and "Remove" takes it out. Level 1 moves
go in the learnset as "db 1, MOVE" too. Level 1 moves of an evolved form are learned as it evolves (the game's
LearnMoveFromLevelUp is called for level 1 too). If moving a starting move
would leave the species with none, its first damaging Gen 9 level 1 move (from
the "Full Gen IX Reference" sheet) that the game has takes its place. Prints what
it did.

Usage: tools/apply_learnset_audit.py <audit.xlsx> <repo root>
"""
import glob, re, sys
import openpyxl

xlsx, repo = sys.argv[1], sys.argv[2]
consts = re.findall(r"^\s*const (\w+)", open(f"{repo}/constants/move_constants.asm").read(), re.M)
MOVES = set(consts[:consts.index("STRUGGLE")])
SPECIAL = {"Psychic": "PSYCHIC_M", "Round": "ROUND_M", "Disarming Voice": "DISARM_VOICE",
           "Stomping Tantrum": "STOMPTANTRUM", "Vise Grip": "VICEGRIP"}


def move_const(name):
    if name in SPECIAL:
        return SPECIAL[name]
    c = re.sub(r"[ \-]", "_", name.upper())
    for cand in (c, c.replace("_", ""), c.replace("_", "", 1)):
        if cand in MOVES:
            return cand
    raise KeyError(name)


evos = open(f"{repo}/data/evos_moves.asm").read()
labels = {re.sub(r"[^a-z]", "", m.lower()): m for m in re.findall(r"^(\w+)EvosMoves:", evos, re.M)}
stats = {}
for path in glob.glob(f"{repo}/data/baseStats/*.asm"):
    stats[re.sub(r"[^a-z]", "", path.split("/")[-1][:-4].lower())] = path


def species_key(name):
    return re.sub(r"[^a-z]", "", name.lower().replace("’", ""))


def learnset_span(text, label):
    start = text.index(f"{label}EvosMoves:")
    ls = text.index("; Learnset\n", start) + len("; Learnset\n")
    end = text.index("\tdb 0\n", ls)
    return ls, end


def starting_moves(path):
    """the text, the starting moves' line spans and the moves (commented-out
    lines between them are kept)"""
    text = open(path).read()
    start = text.index("; attacks known at lvl 0\n") + len("; attacks known at lvl 0\n")
    spans, moves, pos = [], [], start
    while len(moves) < 4:
        end = text.index("\n", pos) + 1
        m = re.match(r"db (\w+)\n", text[pos:end])
        if m:
            spans.append((pos, end))
            moves.append(m.group(1))
        pos = end
    return text, spans, moves


wb = openpyxl.load_workbook(xlsx, read_only=True)
gen9_start = {}
for r in list(wb["Full Gen IX Reference"].iter_rows(values_only=True))[1:]:
    if r[0] and r[5] == 1:
        try:
            gen9_start.setdefault(species_key(r[1]), []).append(move_const(r[3]))
        except KeyError:
            pass
rows = [r for r in list(wb["Recommended Remaster Learnsets"].iter_rows(values_only=True))[1:] if r[0]]
log = []
for r in rows:
    species, level, move, decision = r[1], r[3], move_const(r[4]), r[5]
    key = species_key(species)
    label = labels[key]
    ls, end = learnset_span(evos, label)
    entries = re.findall(r"\tdb (\d+), (\w+)\n", evos[ls:end])
    before = [(int(l), m) for l, m in entries]
    after = [(l, m) for l, m in before if m != move]
    text, sm, start = starting_moves(stats[key])
    new_start = start
    if move in start and level != 1:
        left = [m for m in start if m not in ("0", move)]
        if not left:
            power = dict(re.findall(r"^\tmove (\w+),\s*\w+,\s*(\d+)", open(f"{repo}/data/moves.asm").read(), re.M))
            left = sorted((m for m in gen9_start.get(key, []) if m != move), key=lambda m: power[m] == "0")[:1]
            log.append(f"{species}: {left[0]} starts in {move}'s place (its Gen 9 level 1 move)")
        new_start = left + ["0"] * (4 - len(left))
    if level != "Remove":
        after.append((int(level), move))
    after.sort(key=lambda e: e[0])
    if after != before:
        evos = evos[:ls] + "".join(f"\tdb {l}, {m}\n" for l, m in after) + evos[end:]
    if new_start != start:
        for (a, b), m in reversed(list(zip(sm, new_start))):
            text = text[:a] + f"db {m}\n" + text[b:]
        open(stats[key], "w").write(text)
    if after != before or new_start != start:
        old = [l for l, m in before if m == move] + (["start"] if move in start else [])
        log.append(f"{species}: {move} {','.join(map(str, old)) or '-'} -> {level} ({decision})")
    else:
        log.append(f"{species}: {move} unchanged at {level} ({decision})")
open(f"{repo}/data/evos_moves.asm", "w").write(evos)
print("\n".join(log))
